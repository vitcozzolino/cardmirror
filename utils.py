import csv
import ipaddress
import json
import logging
import re
import socket
import sqlite3
import unicodedata
from io import BytesIO, StringIO
from pathlib import Path
from urllib.parse import urljoin, urlparse

import pandas as pd
import requests
from openpyxl import load_workbook

MAX_REMOTE_REDIRECTS = 5
logger = logging.getLogger(__name__)


class CardOrder:
    def __init__(self, name, set, quantity=1):
        self.name = name
        self.set = set
        self.quantity = quantity

    def __str__(self):
        return f"{self.name}|{self.set}|{self.quantity}"


class CardCandidate:
    def __init__(self, name, card_set, image_url, score):
        self.name = name
        self.card_set = card_set
        self.image_url = image_url
        self.score = score


def load_card_list(uri, header_row=0, rows_to_skip_after_header=0):
    if uri.startswith(("https:", "http:")):
        df = get_csv_from_gcloud(uri, header_row)
    else:
        df = pd.read_csv(uri, header=header_row)

    if rows_to_skip_after_header:
        df = df.iloc[rows_to_skip_after_header:].reset_index(drop=True)

    df.fillna("", inplace=True)
    return df


def load_wanted_cards(uri):
    if uri.lower().endswith(".json"):
        return load_wanted_cards_json(uri)
    return load_card_list(uri)


def load_wanted_cards_json(uri):
    if uri.startswith(("https:", "http:")):
        card_data = json.loads(download_csv(uri))
    else:
        with open(uri, encoding="utf-8") as source:
            card_data = json.load(source)

    if not isinstance(card_data, list):
        raise TypeError("Wanted-card JSON must be an array of card objects")

    cards = []
    for index, card in enumerate(card_data, start=1):
        if not isinstance(card, dict) or not card.get("name"):
            raise ValueError(f"Wanted-card JSON item {index} needs a non-empty 'name'")
        cards.append(
            {
                "Wanted": card["name"],
                "QTY": card.get("quantity", 1),
            }
        )

    return pd.DataFrame(cards)


def load_vendor_price_tiers(uri):
    price_tiers = []
    for row in csv.reader(StringIO(download_csv(uri))):
        if len(row) < 4:
            continue

        quantity_range = row[0].strip()
        price_text = row[3].strip()
        price_match = re.fullmatch(r"\$(\d+(?:\.\d+)?)\s*/\s*card", price_text)
        quantities = re.findall(r"\d+", quantity_range)
        if price_match and quantities:
            price_tiers.append((int(quantities[0]), float(price_match.group(1))))

    if not price_tiers:
        raise ValueError("No single-card price tiers were found in the vendor sheet")

    return sorted(price_tiers)


def load_vendor_catalog(uri, sheet_names, header_row, rows_to_skip_after_header):
    workbook = load_workbook(BytesIO(download_file(uri)), data_only=True)
    header_excel_row = header_row + 1
    first_card_row = header_excel_row + rows_to_skip_after_header + 1
    catalog_columns = {}
    image_links = {}

    for sheet_name in sheet_names:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Vendor workbook does not contain the '{sheet_name}' sheet")

        worksheet = workbook[sheet_name]
        card_set = None
        column_counts = {}
        for column_index in range(1, worksheet.max_column + 1):
            header_value = worksheet.cell(header_excel_row, column_index).value
            if header_value:
                card_set = str(header_value)
            if not card_set:
                continue

            base_column_name = f"{sheet_name}: {card_set}"
            column_counts[base_column_name] = column_counts.get(base_column_name, 0) + 1
            column_name = base_column_name
            if column_counts[base_column_name] > 1:
                column_name = f"{base_column_name} {column_counts[base_column_name]}"

            cards = []
            for row_index in range(first_card_row, worksheet.max_row + 1):
                cell = worksheet.cell(row_index, column_index)
                if not cell.value:
                    continue

                card_name = str(cell.value)
                cards.append(card_name)
                if cell.hyperlink and cell.hyperlink.target:
                    image_links[(column_name, card_name)] = cell.hyperlink.target

            if cards:
                catalog_columns[column_name] = pd.Series(cards)

    if not catalog_columns:
        raise ValueError("No card entries were found in the vendor workbook")

    return pd.DataFrame(catalog_columns), image_links


def sync_vendor_catalog_cache(cache_path, vendor_cards, image_links):
    cache_file = Path(cache_path)
    cache_file.parent.mkdir(parents=True, exist_ok=True)
    card_rows = []
    for card_set in vendor_cards.columns:
        for card_name in vendor_cards[card_set]:
            if not isinstance(card_name, str):
                continue
            card_rows.append(
                (
                    card_name,
                    normalize_catalog_name(card_name),
                    card_set,
                    image_links.get((card_set, card_name)),
                )
            )

    with sqlite3.connect(cache_file) as connection:
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS catalog_cards (
                id INTEGER PRIMARY KEY,
                card_name TEXT NOT NULL,
                normalized_name TEXT NOT NULL,
                card_set TEXT NOT NULL,
                image_url TEXT,
                UNIQUE(card_name, card_set)
            )
            """
        )
        connection.execute("CREATE INDEX IF NOT EXISTS normalized_name_index ON catalog_cards(normalized_name)")
        connection.execute("DELETE FROM catalog_cards")
        connection.executemany(
            """
            INSERT OR IGNORE INTO catalog_cards (
                card_name, normalized_name, card_set, image_url
            ) VALUES (?, ?, ?, ?)
            """,
            card_rows,
        )
        cached_card_count = connection.execute(
            "SELECT COUNT(*) FROM catalog_cards"
        ).fetchone()[0]

    return cached_card_count


def load_vendor_catalog_cache(cache_path):
    with sqlite3.connect(cache_path) as connection:
        rows = connection.execute(
            "SELECT card_name, card_set, image_url FROM catalog_cards ORDER BY id"
        ).fetchall()

    if not rows:
        raise ValueError("Vendor catalog cache is empty")

    catalog_columns = {}
    image_links = {}
    for card_name, card_set, image_url in rows:
        catalog_columns.setdefault(card_set, []).append(card_name)
        if image_url:
            image_links[(card_set, card_name)] = image_url

    return (
        pd.DataFrame(
            {
                card_set: pd.Series(card_names)
                for card_set, card_names in catalog_columns.items()
            }
        ),
        image_links,
    )


def load_vendor_catalog_image_urls(cache_path):
    with sqlite3.connect(cache_path) as connection:
        rows = connection.execute(
            """
            SELECT DISTINCT image_url
            FROM catalog_cards
            WHERE image_url IS NOT NULL AND TRIM(image_url) <> ''
            ORDER BY image_url
            """
        ).fetchall()

    return [image_url for (image_url,) in rows]


def search_vendor_catalog_cache(cache_path, query, limit=30):
    normalized_query = normalize_catalog_name(query)
    if not normalized_query:
        return []

    with sqlite3.connect(cache_path) as connection:
        exact_rows = connection.execute(
            """
            SELECT catalog_cards.id, catalog_cards.card_name,
                   catalog_cards.normalized_name, catalog_cards.card_set,
                   catalog_cards.image_url
            FROM catalog_cards
            INNER JOIN (
                SELECT normalized_name, MIN(id) AS id
                FROM catalog_cards
                WHERE normalized_name = ?
                   OR instr(' ' || normalized_name || ' ', ' ' || ? || ' ') > 0
                GROUP BY normalized_name
            ) AS first_matches ON catalog_cards.id = first_matches.id
            ORDER BY catalog_cards.normalized_name, card_set
            LIMIT ?
            """,
            (normalized_query, normalized_query, limit),
        ).fetchall()
        rows = exact_rows or connection.execute(
            """
            SELECT catalog_cards.id, catalog_cards.card_name,
                   catalog_cards.normalized_name, catalog_cards.card_set,
                   catalog_cards.image_url
            FROM catalog_cards
            INNER JOIN (
                SELECT normalized_name, MIN(id) AS id
                FROM catalog_cards
                WHERE normalized_name LIKE ?
                GROUP BY normalized_name
            ) AS first_matches ON catalog_cards.id = first_matches.id
            ORDER BY catalog_cards.normalized_name, card_set
            LIMIT ?
            """,
            (f"%{normalized_query}%", limit),
        ).fetchall()

    return [
        {
            "id": card_id,
            "name": card_name,
            "normalized_name": normalized_name,
            "card_set": card_set,
            "image_url": image_url,
        }
        for card_id, card_name, normalized_name, card_set, image_url in rows
    ]


def search_vendor_catalog_versions(cache_path, query, limit=200):
    normalized_query = normalize_catalog_name(query)
    if not normalized_query:
        return []

    with sqlite3.connect(cache_path) as connection:
        exact_rows = connection.execute(
            """
            SELECT id, card_name, normalized_name, card_set, image_url
            FROM catalog_cards
            WHERE normalized_name = ?
               OR instr(' ' || normalized_name || ' ', ' ' || ? || ' ') > 0
            ORDER BY card_set, card_name, id
            LIMIT ?
            """,
            (normalized_query, normalized_query, limit),
        ).fetchall()
        rows = exact_rows or connection.execute(
            """
            SELECT id, card_name, normalized_name, card_set, image_url
            FROM catalog_cards
            WHERE normalized_name LIKE ?
            ORDER BY card_set, card_name, id
            LIMIT ?
            """,
            (f"%{normalized_query}%", limit),
        ).fetchall()

    return [
        {
            "id": card_id,
            "name": card_name,
            "normalized_name": version_normalized_name,
            "card_set": card_set,
            "image_url": image_url,
        }
        for card_id, card_name, version_normalized_name, card_set, image_url in rows
    ]


def load_vendor_catalog_card(cache_path, card_id):
    with sqlite3.connect(cache_path) as connection:
        row = connection.execute(
            """
            SELECT card_name, card_set, image_url
            FROM catalog_cards
            WHERE id = ?
            """,
            (card_id,),
        ).fetchone()

    if not row:
        return None
    return CardCandidate(row[0], row[1], row[2], 1.0)


def normalize_catalog_name(card_name):
    normalized = unicodedata.normalize("NFKD", card_name)
    normalized = normalized.replace("Æ", "AE").replace("æ", "ae")
    normalized = normalized.encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", normalized).strip()


def get_csv_from_gcloud(uri, header_row=0):
    return pd.read_csv(StringIO(download_csv(uri)), header=header_row)


def download_csv(uri):
    logger.info("Downloading CSV list from %s...", uri)
    response = request_remote(uri, timeout=30)
    try:
        response.raise_for_status()
        return response.text
    finally:
        response.close()


def download_file(uri):
    logger.info("Downloading file from %s...", uri)
    response = request_remote(uri, timeout=30)
    try:
        response.raise_for_status()
        return response.content
    finally:
        response.close()


def request_remote(uri, timeout=30, **kwargs):
    """Fetch a public HTTP(S) URL without following unsafe redirects."""
    current_uri = uri
    for _ in range(MAX_REMOTE_REDIRECTS + 1):
        validate_remote_url(current_uri)
        response = requests.get(
            current_uri,
            allow_redirects=False,
            timeout=timeout,
            **kwargs,
        )
        if not response.is_redirect:
            return response

        location = response.headers.get("Location")
        response.close()
        if not location:
            raise ValueError("Remote response contained an empty redirect")
        current_uri = urljoin(current_uri, location)

    raise ValueError("Remote URL exceeded the redirect limit")


def validate_remote_url(uri):
    parsed_uri = urlparse(uri)
    if parsed_uri.scheme not in {"http", "https"}:
        raise ValueError("Remote URL must use HTTP or HTTPS")
    if parsed_uri.username or parsed_uri.password or not parsed_uri.hostname:
        raise ValueError("Remote URL must not contain credentials or an empty host")

    try:
        addresses = {
            address[4][0]
            for address in socket.getaddrinfo(
                parsed_uri.hostname,
                parsed_uri.port,
                type=socket.SOCK_STREAM,
            )
        }
    except (OSError, ValueError) as error:
        raise ValueError(f"Could not resolve remote host: {parsed_uri.hostname}") from error

    for address in addresses:
        parsed_address = ipaddress.ip_address(address)
        if (
            parsed_address.is_private
            or parsed_address.is_loopback
            or parsed_address.is_link_local
            or parsed_address.is_reserved
            or parsed_address.is_multicast
            or parsed_address.is_unspecified
        ):
            raise ValueError("Remote URL resolves to a non-public address")
